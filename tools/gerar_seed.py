"""Le as planilhas *_empresas_emails_telefones.xlsx e gera db/seed/cadastros.csv.

Uso: python tools/gerar_seed.py   (requer: pip install openpyxl)
"""
import csv
import glob
import os
import re

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "db", "seed", "cadastros.csv")
CNPJ_INTERNO = "00000000000000"  # linhas da propria Avanco, nao sao clientes


def limpo(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def digitos(v):
    return re.sub(r"\D", "", str(v)) if v is not None else ""


linhas, vistos = [], set()
ignoradas = {"sem_email": 0, "cnpj_interno": 0, "sem_cnpj_e_sem_conta": 0, "duplicada": 0}
origens_desconhecidas = 0

for arq in sorted(glob.glob(os.path.join(RAIZ, "*_empresas_emails_telefones.xlsx"))):
    base = os.path.basename(arq).split("_")[0]
    ws = openpyxl.load_workbook(arq, read_only=True).active
    cab = [limpo(c) for c in next(ws.iter_rows(values_only=True))]
    col = {}
    for i, nome in enumerate(cab):
        n = nome.lower()
        if "codigo_fortes" in n: col["conta"] = i
        elif n == "cnpj": col["cnpj"] = i
        elif n.startswith("raz"): col["razao"] = i
        elif n == "telefone": col["tel"] = i
        elif n == "email": col["email"] = i
        elif n == "origem": col["origem"] = i
        elif n == "empresa nome": col["fantasia"] = i

    for r in ws.iter_rows(min_row=2, values_only=True):
        email = limpo(r[col["email"]]).lower()
        cnpj = digitos(r[col["cnpj"]])
        origem = limpo(r[col["origem"]])
        if not email or "@" not in email:
            ignoradas["sem_email"] += 1; continue
        conta = limpo(r[col["conta"]])
        if not cnpj and not conta:
            ignoradas["sem_cnpj_e_sem_conta"] += 1; continue
        if cnpj == CNPJ_INTERNO:
            ignoradas["cnpj_interno"] += 1; continue

        m = re.match(r"^usuario\s*\((.*)\)\s*$", origem, re.I)
        if m:
            tipo, nome = "usuario", m.group(1).strip()
        else:
            tipo, nome = "filial", ""
            if origem.lower() != "filial":
                origens_desconhecidas += 1

        conta = limpo(r[col["conta"]])
        chave = (base, tipo, email, cnpj, conta)
        if chave in vistos:
            ignoradas["duplicada"] += 1; continue
        vistos.add(chave)
        linhas.append([base, tipo, nome, email, cnpj, conta, limpo(r[col["tel"]]),
                       limpo(r[col["razao"]]), limpo(r[col["fantasia"]])])

with open(SAIDA, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["base", "tipo", "nome_usuario", "email", "cnpj", "conta_fortes",
                "telefone", "razao_social", "nome_fantasia"])
    w.writerows(linhas)

print(f"{len(linhas)} registros gravados em {SAIDA}")
print("usuarios:", sum(1 for l in linhas if l[1] == "usuario"),
      "| filiais:", sum(1 for l in linhas if l[1] == "filial"))
print("ignoradas:", ignoradas, "| origens fora do padrao tratadas como filial:", origens_desconhecidas)
